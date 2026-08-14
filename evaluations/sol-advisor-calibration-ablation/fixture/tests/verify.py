#!/usr/bin/env python3
"""Deterministic oracle for the synthetic WGCNA repair."""

from __future__ import annotations

import csv
import math
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "tests" / "data"
PIPELINE = ROOT / "src" / "pipeline.py"
EXPECTED_OUTPUTS = {
    "module_assignments.tsv",
    "module_eigengenes.tsv",
    "module_trait_associations.tsv",
    "status.tsv",
    "wgcna-report.xlsx",
}


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8", newline="") as stream:
        return list(csv.DictReader(stream, delimiter="\t"))


def run_pipeline(
    output: Path, *, traits: str, modules: str
) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [
            sys.executable,
            str(PIPELINE),
            "--expression",
            str(DATA / "expression.tsv"),
            "--traits",
            str(DATA / traits),
            "--modules",
            str(DATA / modules),
            "--output",
            str(output),
        ],
        cwd=ROOT,
        check=False,
        capture_output=True,
        text=True,
    )


class WgcnaContract(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls._temporary = tempfile.TemporaryDirectory(prefix="calibration-wgcna-")
        temporary = Path(cls._temporary.name)
        cls.normal = temporary / "normal"
        cls.grey = temporary / "grey"
        cls.mismatch = temporary / "mismatch"
        cls.normal_result = run_pipeline(
            cls.normal, traits="traits.tsv", modules="modules.tsv"
        )
        cls.grey_result = run_pipeline(
            cls.grey, traits="traits.tsv", modules="modules-grey.tsv"
        )
        cls.mismatch_result = run_pipeline(
            cls.mismatch,
            traits="traits-mismatch.tsv",
            modules="modules.tsv",
        )

    @classmethod
    def tearDownClass(cls) -> None:
        cls._temporary.cleanup()

    def require_normal(self) -> None:
        self.assertEqual(self.normal_result.returncode, 0, self.normal_result.stderr)

    def test_01_nonnumeric_sample_ids_survive_r_worker(self) -> None:
        self.require_normal()
        rows = read_tsv(self.normal / "module_eigengenes.tsv")
        self.assertEqual(
            [row["SampleId"] for row in rows], ["S-01", "S-02", "S-03", "S-04"]
        )

    def test_02_two_level_trait_has_finite_associations(self) -> None:
        self.require_normal()
        rows = read_tsv(self.normal / "module_trait_associations.tsv")
        group_rows = [row for row in rows if row["Trait"] == "Group"]
        self.assertEqual({row["Module"] for row in group_rows}, {"blue", "brown"})
        self.assertTrue(
            all(math.isfinite(float(row["Correlation"])) for row in group_rows)
        )

    def test_03_mismatch_is_field_specific(self) -> None:
        self.assertNotEqual(self.mismatch_result.returncode, 0)
        self.assertIn(
            "SampleId mismatch between expression and traits: "
            "missing_in_traits=S-04; extra_in_traits=S-99",
            self.mismatch_result.stderr,
        )
        self.assertNotIn("Traceback", self.mismatch_result.stderr)

    def test_04_grey_only_emits_structured_empty_outputs(self) -> None:
        self.assertEqual(self.grey_result.returncode, 0, self.grey_result.stderr)
        self.assertEqual({path.name for path in self.grey.iterdir()}, EXPECTED_OUTPUTS)
        self.assertEqual(
            read_tsv(self.grey / "status.tsv"),
            [{"status": "no_modules", "message": "only grey genes remain"}],
        )
        self.assertEqual(read_tsv(self.grey / "module_eigengenes.tsv"), [])
        self.assertEqual(read_tsv(self.grey / "module_trait_associations.tsv"), [])
        self.assertNotIn("Traceback", self.grey_result.stderr)

    def test_05_workbook_sheet_contract(self) -> None:
        self.require_normal()
        workbook = load_workbook(
            self.normal / "wgcna-report.xlsx", read_only=True, data_only=True
        )
        self.assertEqual(
            workbook.sheetnames,
            ["Module Eigengenes", "Module Assignments", "Trait Associations", "README"],
        )

    def test_06_workbook_preserves_sample_ids(self) -> None:
        self.require_normal()
        workbook = load_workbook(
            self.normal / "wgcna-report.xlsx", read_only=True, data_only=True
        )
        values = list(workbook["Module Eigengenes"].values)
        self.assertEqual(
            [row[0] for row in values[1:]], ["S-01", "S-02", "S-03", "S-04"]
        )

    def test_07_workbook_preserves_module_assignments(self) -> None:
        self.require_normal()
        workbook = load_workbook(
            self.normal / "wgcna-report.xlsx", read_only=True, data_only=True
        )
        values = list(workbook["Module Assignments"].values)
        self.assertEqual(
            values[1:],
            [
                ("GeneA", "blue"),
                ("GeneB", "blue"),
                ("GeneC", "brown"),
                ("GeneD", "brown"),
            ],
        )

    def test_08_workbook_preserves_association_shape(self) -> None:
        self.require_normal()
        workbook = load_workbook(
            self.normal / "wgcna-report.xlsx", read_only=True, data_only=True
        )
        values = list(workbook["Trait Associations"].values)
        self.assertEqual(values[0], ("Module", "Trait", "Correlation", "N"))
        self.assertEqual(len(values) - 1, 4)

    def test_09_workbook_readme_records_status_and_contract(self) -> None:
        self.require_normal()
        workbook = load_workbook(
            self.normal / "wgcna-report.xlsx", read_only=True, data_only=True
        )
        values = list(workbook["README"].values)
        self.assertEqual(values[1], ("Status", "ok"))
        self.assertEqual(
            values[2], ("Contract", "SampleId-preserving module/trait output")
        )
        self.assertEqual(
            {path.name for path in self.normal.iterdir()}, EXPECTED_OUTPUTS
        )


if __name__ == "__main__":
    unittest.main(verbosity=2)
