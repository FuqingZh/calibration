"""Stable workbook materialization for the synthetic fixture."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook


def _rows(path: Path) -> list[list[str]]:
    with path.open(encoding="utf-8", newline="") as stream:
        return list(csv.reader(stream, delimiter="\t"))


def export_workbook(output_dir: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    sources = (
        ("Module Eigengenes", "module_eigengenes.tsv"),
        ("Module Assignments", "module_assignments.tsv"),
        ("Trait Associations", "module_trait_associations.tsv"),
    )
    for sheet_name, filename in sources:
        sheet = workbook.create_sheet(sheet_name)
        for row in _rows(output_dir / filename):
            sheet.append(row)
    readme = workbook.create_sheet("README")
    status = _rows(output_dir / "status.tsv")
    readme.append(["Synthetic WGCNA report"])
    readme.append(["Status", status[1][0]])
    readme.append(["Contract", "SampleId-preserving module/trait output"])
    workbook.save(output_dir / "wgcna-report.xlsx")
